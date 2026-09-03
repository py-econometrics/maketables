"""Tests for embedded line breaks (\\n) in labels, headers, and notes.

Data cells already supported this before this change; these tests cover the
previously-missing positions: row labels, column headers/spanners, and
notes, across the tex, typst, and gt/html backends. Docx already handled
all of these implicitly (python-docx converts embedded newlines on its own),
so it isn't covered here. Captions do not support line breaks - a caption
is a single-line label by design.

TestHeaderVerticalCentering covers a related feature: when a header row
mixes cells with different line counts (e.g. one model_heads entry has a
line break and its sibling doesn't), the shorter cells should be vertically
centered against the tallest one, in tex, typst, and docx (not html/gt,
where this is great_tables' own default rendering behavior).
"""

from helpers import normalize_html

import maketables as mt


class TestLineBreaksTex:
    """Line breaks should render as a real \\\\ inside \\makecell{}."""

    def test_label_linebreak(self, fitted_model, snapshot):
        table = mt.ETable([fitted_model], labels={"x": "Line one\nLine two"})
        assert table.make(type="tex") == snapshot

    def test_header_linebreak(self, fitted_model, snapshot):
        table = mt.ETable([fitted_model], model_heads=["Line one\nLine two"])
        assert table.make(type="tex") == snapshot

    def test_notes_linebreak(self, fitted_model, snapshot):
        table = mt.ETable([fitted_model], notes="Line one\nLine two")
        assert table.make(type="tex") == snapshot


class TestLineBreaksGt:
    """Line breaks should render as a real <br>, not escaped &lt;br&gt; text."""

    def test_label_linebreak(self, fitted_model, snapshot):
        table = mt.ETable([fitted_model], labels={"x": "Line one\nLine two"})
        assert normalize_html(table.make(type="gt").as_raw_html()) == snapshot

    def test_header_linebreak(self, fitted_model, snapshot):
        table = mt.ETable([fitted_model], model_heads=["Line one\nLine two"])
        assert normalize_html(table.make(type="gt").as_raw_html()) == snapshot

    def test_notes_linebreak(self, fitted_model, snapshot):
        table = mt.ETable([fitted_model], notes="Line one\nLine two")
        assert normalize_html(table.make(type="gt").as_raw_html()) == snapshot

    def test_special_characters_still_escaped_alongside_linebreak(self, fitted_model):
        """A real <br> must appear, but any other special characters on
        either side of it must still be escaped, not passed through raw.
        """
        table = mt.ETable([fitted_model], notes="A<B\nC>D")
        html = table.make(type="gt").as_raw_html()
        assert "A&lt;B<br>C&gt;D" in html
        assert "<br>" in html


class TestLineBreaksTypst:
    """Line breaks should render as Typst's native ' \\ ' line-break syntax."""

    def test_label_linebreak(self, fitted_model):
        table = mt.ETable([fitted_model], labels={"x": "Line one\nLine two"})
        typst = table.make(type="typst")
        assert "Line one \\ Line two" in typst

    def test_header_linebreak(self, fitted_model):
        table = mt.ETable([fitted_model], model_heads=["Line one\nLine two"])
        typst = table.make(type="typst")
        assert "Line one \\ Line two" in typst

    def test_notes_linebreak(self, fitted_model):
        table = mt.ETable([fitted_model], notes="Line one\nLine two")
        typst = table.make(type="typst")
        assert "Line one \\ Line two" in typst

    def test_escaping_still_applied_alongside_linebreak(self, fitted_model):
        """Each line must still be escaped individually (e.g. a literal
        underscore), not just naively joined.
        """
        table = mt.ETable([fitted_model], labels={"x": "a_1\nb_2"})
        typst = table.make(type="typst")
        assert r"a\_1 \ b\_2" in typst


class TestHeaderVerticalCentering:
    """A header row mixing single-line and multi-line cells should center
    the shorter cells against the tallest sibling, not leave them flush at
    the top.
    """

    def test_tex_shorter_header_gets_raisebox(self, fitted_models):
        table = mt.ETable(
            fitted_models, model_heads=["Simple", "With Fixed\nEffects"]
        )
        tex = table.make(type="tex")
        assert r"\raisebox{-0.5\baselineskip}{Simple}" in tex
        assert r"\makecell{With Fixed\\Effects}" in tex

    def test_tex_equal_height_row_unaffected(self, fitted_models):
        table = mt.ETable(fitted_models, model_heads=["Simple", "Simple"])
        tex = table.make(type="tex")
        assert "raisebox" not in tex

    def test_typst_header_cells_get_horizon_align(self, fitted_models):
        table = mt.ETable(
            fitted_models, model_heads=["Simple", "With Fixed\nEffects"]
        )
        typst = table.make(type="typst")
        assert "#table.cell(colspan: 1, align: horizon)[Simple]" in typst
        assert (
            "#table.cell(colspan: 1, align: horizon)[With Fixed \\ Effects]" in typst
        )

    def test_docx_header_rows_vertically_centered(self, fitted_models):
        from docx.enum.table import WD_ALIGN_VERTICAL

        table = mt.ETable(
            fitted_models, model_heads=["Simple", "With Fixed\nEffects"]
        )
        doc = table.make(type="docx")
        table_obj = doc.tables[0]
        model_heads_row = table_obj.rows[1]
        assert all(
            cell.vertical_alignment == WD_ALIGN_VERTICAL.CENTER
            for cell in model_heads_row.cells
        )
        data_row = table_obj.rows[3]  # first coefficient row ("x")
        assert all(cell.vertical_alignment is None for cell in data_row.cells)
