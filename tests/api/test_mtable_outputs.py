"""Tests for MTable: captions, notes, row groups, save/update and display."""
# ruff: noqa: SLF001

import json
from unittest import mock

import pandas as pd
import pytest
from docx import Document

import maketables as mt
from maketables import MTable


def _frame(*, grouped=False, spanners=False):
    cols = ["(1)", "(2)", "(3)"]
    if spanners:
        cols = pd.MultiIndex.from_tuples(
            [("Wage", "(1)"), ("Wage", "(2)"), ("Hours", "(3)")]
        )
    index = ["alpha", "beta", "gamma", "delta"]
    if grouped:
        index = pd.MultiIndex.from_tuples(
            [("Coefs", "alpha"), ("Coefs", "beta"), ("Stats", "gamma"), ("Stats", "d")]
        )
    return pd.DataFrame(
        [
            ["1.0", "2.0", "3.0"],
            ["0.1", "0.2", "0.3"],
            ["10", "20", "30"],
            ["a", "b", "c"],
        ],
        index=index,
        columns=cols,
    )


def _table(*, grouped=False, spanners=False, **kwargs):
    kwargs.setdefault("caption", "My caption")
    kwargs.setdefault("tab_label", "tab:mine")
    kwargs.setdefault("notes", "First note.\nSecond note.")
    return MTable(_frame(grouped=grouped, spanners=spanners), **kwargs)


@pytest.mark.parametrize("fmt", ["tex", "typst"])
@pytest.mark.parametrize("grouped", [False, True])
@pytest.mark.parametrize("spanners", [False, True])
def test_text_formats_contain_caption_notes_and_cells(fmt, grouped, spanners):
    out = _table(grouped=grouped, spanners=spanners).make(type=fmt)
    assert "My caption" in out
    assert "First note." in out
    assert "Second note." in out
    assert "tab:mine" in out
    assert "alpha" in out
    if grouped:
        assert "Coefs" in out
        assert "Stats" in out
    if spanners:
        assert "Wage" in out
        assert "Hours" in out


@pytest.mark.parametrize("fmt", ["tex", "typst"])
def test_rgroup_display_off_hides_group_names(fmt):
    out = _table(grouped=True, rgroup_display=False).make(type=fmt)
    assert "Coefs" not in out
    assert "Stats" not in out


@pytest.mark.parametrize("sep", ["tb", "t", "b", ""])
@pytest.mark.parametrize("fmt", ["tex", "typst", "gt", "docx"])
def test_rgroup_sep_variants_render(fmt, sep):
    out = _table(grouped=True, spanners=True, rgroup_sep=sep).make(type=fmt)
    assert out is not None


def test_no_caption_no_notes():
    t = _table(caption=None, notes="", tab_label=None)
    assert "\\caption" not in t.make(type="tex")
    assert "caption" not in t.make(type="typst")
    doc = t.make(type="docx")
    assert not any(p.style.name == "Caption" for p in doc.paragraphs)


def test_gt_output_has_caption_notes_and_groups():
    html = _table(grouped=True, spanners=True).make(type="html")
    assert "My caption" in html
    assert "First note." in html
    assert "Second note." in html
    assert "Coefs" in html
    assert "Wage" in html


def test_gt_style_options():
    t = _table(grouped=True, spanners=True)
    gt = t.make(
        type="gt",
        gt_style={
            "align": "center",
            "table_width": "80%",
            "first_col_width": "150px",
            "table_font_size_all": "12px",
        },
    )
    html = gt.as_raw_html()
    assert "80%" in html
    assert "150px" in html


def test_tex_style_options():
    out = _table(grouped=True, spanners=True).make(
        type="tex",
        tex_style={
            "tab_width": r"0.8\textwidth",
            "first_col_width": "3cm",
            "texlocation": "htbp",
            "data_align": "c",
        },
    )
    assert "tabularx" in out
    assert "3cm" in out
    assert "htbp" in out
    out = _table().make(type="tex", tex_style={"tab_width": "\\linewidth"})
    assert "tabularx" in out


def test_typst_style_options():
    out = _table(grouped=True, spanners=True).make(
        type="typst",
        typst_style={"first_col_width": "2.5cm", "notes_font_size": "9pt"},
    )
    assert "2.5cm" in out
    assert "9pt" in out


def test_special_characters_in_typst_and_docx():
    df = pd.DataFrame({"a_b": ["50% & $x$ #1 _y_"]}, index=["r_1"])
    t = MTable(df, caption="Cap & 100%", notes="Note #1 & more")
    assert "r\\_1" in t.make(type="typst")
    assert "\\#" in t.make(type="typst")
    assert t.make(type="docx").tables[0].cell(1, 1).text


def test_linebreaks_in_cells():
    df = pd.DataFrame({"A": ["x<br>y"]}, index=["row<br>two"])
    t = MTable(df, notes="n1<br>n2")
    for fmt in ("tex", "typst", "gt", "docx"):
        assert t.make(type=fmt) is not None


def test_empty_cells_and_nan():
    df = pd.DataFrame({"A": ["", None], "B": ["1", "2"]}, index=["a", "b"])
    t = MTable(df)
    for fmt in ("tex", "typst", "gt", "docx"):
        assert t.make(type=fmt) is not None


def test_constructor_validation():
    with pytest.raises(TypeError):
        MTable([[1, 2]])
    idx = pd.MultiIndex.from_tuples([("a", "b", "c")])
    with pytest.raises(ValueError, match="at most two levels"):
        MTable(pd.DataFrame({"x": [1]}, index=idx))


def test_make_rejects_unknown_type():
    with pytest.raises(ValueError, match="types must be"):
        _table().make(type="pdf")


def test_call_is_make():
    t = _table()
    assert t(type="tex") == t.make(type="tex")
    assert repr(t) == ""


def test_default_paths_forms(tmp_path):
    assert _table(default_paths=str(tmp_path)).default_paths["tex"] == str(tmp_path)
    d = {"tex": str(tmp_path)}
    assert _table(default_paths=d).default_paths == d
    assert _table(default_paths=None).default_paths == {}


# ---------------------------------------------------------------- docx


def test_docx_caption_notes_and_table():
    doc = _table(grouped=True, spanners=True).make(type="docx")
    texts = [p.text for p in doc.paragraphs]
    assert any("My caption" in x for x in texts)
    table = doc.tables[0]
    cell_text = " ".join(c.text for row in table.rows for c in row.cells)
    assert "Coefs" in cell_text
    assert "Wage" in cell_text
    assert "alpha" in cell_text
    assert "First note." in cell_text


def test_docx_style_options():
    doc = _table(grouped=True).make(
        type="docx",
        docx_style={
            "first_col_width": "3cm",
            "font_name": "Arial",
            "font_size_pt": 9,
            "caption_align": "center",
            "notes_align": "right",
            "align_center_cells": False,
        },
    )
    assert doc.tables
    for align in ("left", "right"):
        _table().make(type="docx", docx_style={"caption_align": align})


# ---------------------------------------------------------------- save


@pytest.mark.parametrize("type_", ["tex", "typst", "html", "docx"])
def test_save_writes_file_with_extension_added(tmp_path, type_):
    t = _table(grouped=True)
    result = t.save(type=type_, file_name=str(tmp_path / "out"))
    assert result is not None  # show=True returns GT
    path = tmp_path / ("out." + type_)
    assert path.exists()
    if type_ == "docx":
        assert Document(str(path)).tables
    elif type_ in ("tex", "typst"):
        assert path.read_text() == t.make(type=type_)
    else:
        assert "My caption" in path.read_text()


def test_save_show_false_returns_none(tmp_path):
    assert _table().save("tex", str(tmp_path / "a.tex"), show=False) is None


def test_save_uses_default_path_and_tab_label(tmp_path):
    t = _table(default_paths=str(tmp_path) + "/", tab_label="mylabel.tex")
    t.save("tex", show=False)
    assert (tmp_path / "mylabel.tex").exists()


def test_save_relative_name_joins_default_path(tmp_path):
    t = _table(default_paths={"tex": str(tmp_path)})
    t.save("tex", "rel.tex", show=False)
    assert (tmp_path / "rel.tex").exists()


def test_save_errors(tmp_path):
    with pytest.raises(ValueError, match="types must be"):
        _table().save(type="pdf", file_name="x")
    with pytest.raises(ValueError, match="tab_label"):
        _table(tab_label=None).save("tex")
    with pytest.raises(ValueError, match="Default path"):
        _table(default_paths=None).save("tex")
    target = tmp_path / "x.tex"
    target.write_text("old")
    with pytest.raises(ValueError, match="already exists"):
        _table().save("tex", str(target), replace=False)
    _table().save("tex", str(target), replace=True, show=False)
    assert target.read_text() != "old"
    with pytest.raises(ValueError, match="not a valid path"):
        _table().save("tex", str(tmp_path / "missing" / "x.tex"))


# ---------------------------------------------------------------- update


def test_update_tex_creates_appends_and_replaces(tmp_path):
    f = tmp_path / "doc.tex"
    _table(tab_label="tab:one").update_tex(str(f))
    assert "\\label{tab:one}" in f.read_text()
    _table(tab_label="tab:two", caption="Second").update_tex(str(f))
    both = f.read_text()
    assert "tab:one" in both
    assert "tab:two" in both
    _table(tab_label="tab:one", caption="Replaced").update_tex(str(f))
    again = f.read_text()
    assert again.count("\\label{tab:one}") == 1
    assert "Replaced" in again
    assert "tab:two" in again


def test_update_tex_inserts_before_end_document(tmp_path):
    f = tmp_path / "doc.tex"
    f.write_text("\\begin{document}\nHello\n\\end{document}\n")
    _table().update_tex(str(f))
    text = f.read_text()
    assert text.index("\\label{tab:mine}") < text.index("\\end{document}")
    f.write_text("\\begin{document}\nHello\n\n\\end{document}\n")
    _table(tab_label="tab:b").update_tex(str(f))
    assert "tab:b" in f.read_text()
    f.write_text("\\begin{document}\nHello\\end{document}\n")
    _table(tab_label="tab:c").update_tex(str(f))
    assert "tab:c" in f.read_text()


def test_update_tex_appends_with_spacing(tmp_path):
    f = tmp_path / "doc.tex"
    f.write_text("text")
    _table().update_tex(str(f))
    assert f.read_text().startswith("text\n\n")
    f.write_text("text\n")
    _table(tab_label="tab:b").update_tex(str(f))
    assert f.read_text().startswith("text\n\n")


def test_update_tex_defaults_and_errors(tmp_path):
    t = _table(default_paths=str(tmp_path), tab_label="lbl")
    t.update_tex()
    assert (tmp_path / "lbl.tex").exists()
    t.update_tex("rel", tab_label="tab:x", tex_style={"texlocation": "h"})
    assert (tmp_path / "rel.tex").exists()
    assert t.tab_label == "lbl"
    assert t.update_tex("rel2.tex", show=True) is not None
    with pytest.raises(ValueError, match="tab_label"):
        _table(tab_label=None).update_tex(str(tmp_path / "a.tex"))
    with pytest.raises(ValueError, match="Default path"):
        _table(default_paths=None).update_tex()
    with pytest.raises(ValueError, match=r"\.tex extension"):
        _table().update_tex(str(tmp_path / "a.txt"))
    with pytest.raises(ValueError, match="not a valid path"):
        _table().update_tex(str(tmp_path / "nope" / "a.tex"))


def test_update_typst_creates_appends_and_replaces(tmp_path):
    f = tmp_path / "doc.typ"
    _table(tab_label="tab:one").update_typst(str(f))
    assert "<tab:one>" in f.read_text()
    _table(tab_label="tab:two", caption="Second").update_typst(str(f))
    both = f.read_text()
    assert "<tab:one>" in both
    assert "<tab:two>" in both
    _table(tab_label="tab:one", caption="Replaced").update_typst(str(f))
    again = f.read_text()
    assert again.count("<tab:one>") == 1
    assert "Replaced" in again
    assert "<tab:two>" in again


def test_update_typst_appends_to_existing_text(tmp_path):
    f = tmp_path / "doc.typ"
    f.write_text("= Heading")
    _table().update_typst(str(f))
    assert f.read_text().startswith("= Heading\n\n")
    f.write_text("= Heading\n")
    _table(tab_label="tab:b").update_typst(str(f))
    assert f.read_text().startswith("= Heading\n\n")


def test_update_typst_defaults_and_errors(tmp_path):
    t = _table(default_paths=str(tmp_path), tab_label="lbl")
    t.update_typst()
    assert (tmp_path / "lbl.typ").exists()
    t.update_typst(
        "rel.typst", tab_label="tab:x", typst_style={"notes_font_size": "8pt"}
    )
    assert (tmp_path / "rel.typst").exists()
    assert t.update_typst("rel2", show=True) is not None
    with pytest.raises(ValueError, match="tab_label"):
        _table(tab_label=None).update_typst(str(tmp_path / "a.typ"))
    with pytest.raises(ValueError, match="Default path"):
        _table(default_paths=None).update_typst()
    with pytest.raises(ValueError, match="extension"):
        _table().update_typst(str(tmp_path / "a.txt"))
    with pytest.raises(ValueError, match="not a valid path"):
        _table().update_typst(str(tmp_path / "nope" / "a.typ"))


def test_update_docx_creates_appends_and_replaces(tmp_path):
    f = tmp_path / "doc.docx"
    _table(caption="First").update_docx(str(f))
    assert len(Document(str(f)).tables) == 1
    _table(caption="Second").update_docx(str(f))
    assert len(Document(str(f)).tables) == 2
    _table(caption="Replaced").update_docx(str(f), tab_num=1)
    doc = Document(str(f))
    assert len(doc.tables) == 2
    assert any("Replaced" in p.text for p in doc.paragraphs)
    assert not any("First" in p.text for p in doc.paragraphs)
    # out-of-range tab_num appends
    _table().update_docx(str(f), tab_num=9)
    assert len(Document(str(f)).tables) == 3


def test_update_docx_replace_without_caption(tmp_path):
    f = tmp_path / "doc.docx"
    _table(caption=None).update_docx(str(f))
    _table(caption=None, notes="").update_docx(str(f), tab_num=1)
    assert len(Document(str(f)).tables) == 1


def test_update_docx_defaults_and_errors(tmp_path):
    t = _table(default_paths=str(tmp_path))
    t.update_docx("rel", docx_style={"font_size_pt": 8})
    assert (tmp_path / "rel.docx").exists()
    assert t.update_docx("rel2.docx", show=True) is not None
    with pytest.raises(ValueError, match="file_name must be provided"):
        t.update_docx()
    with pytest.raises(ValueError, match=r"\.docx extension"):
        _table().update_docx(str(tmp_path / "a.txt"))
    with pytest.raises(ValueError, match="not a valid path"):
        _table().update_docx(str(tmp_path / "nope" / "a.docx"))


# ---------------------------------------------------------------- display


def _write_info(tmp_path, monkeypatch, base, target=None):
    ident = {"base-format": base}
    if target is not None:
        ident["target-format"] = target
    info = tmp_path / "info.json"
    info.write_text(json.dumps({"format": {"identifier": ident}}), encoding="utf-8")
    monkeypatch.setenv("QUARTO_EXECUTE_INFO", str(info))


def test_quarto_without_context_falls_back_to_gt(monkeypatch):
    monkeypatch.delenv("QUARTO_EXECUTE_INFO", raising=False)
    t = _table()
    assert t._resolve_quarto_output_type() == "gt"
    assert type(t.make(type="quarto")).__name__ == "GT"


def test_quarto_missing_or_invalid_info_falls_back(tmp_path, monkeypatch):
    t = _table()
    monkeypatch.setenv("QUARTO_EXECUTE_INFO", str(tmp_path / "missing.json"))
    assert t._resolve_quarto_output_type() == "gt"
    bad = tmp_path / "bad.json"
    bad.write_text("{not json", encoding="utf-8")
    monkeypatch.setenv("QUARTO_EXECUTE_INFO", str(bad))
    assert t._resolve_quarto_output_type() == "gt"
    monkeypatch.setenv("QUARTO_EXECUTE_INFO", str(tmp_path))  # a directory
    assert t._resolve_quarto_output_type() == "gt"


@pytest.mark.parametrize(
    ("base", "target", "expected"),
    [
        ("html", None, "html"),
        ("html", "html", "html"),
        ("latex", "latex", "tex"),
        ("pdf", None, "tex"),
        ("latex", "typst", "typst"),
        ("typst", None, "typst"),
        ("docx", None, "docx"),
        ("epub", None, "gt"),
        ("epub", "docx", "docx"),
    ],
)
def test_quarto_resolution(tmp_path, monkeypatch, base, target, expected):
    _write_info(tmp_path, monkeypatch, base, target)
    assert _table()._resolve_quarto_output_type() == expected


def test_quarto_make_per_target(tmp_path, monkeypatch):
    t = _table()
    _write_info(tmp_path, monkeypatch, "html")
    html = t.make(type="quarto")
    assert isinstance(html, str)
    assert "<table" in html

    _write_info(tmp_path, monkeypatch, "latex")
    tex = t.make(type="quarto")
    assert str(tex) == t.make(type="tex")
    bundle = tex._repr_mimebundle_()
    assert bundle["text/latex"] == str(tex)
    assert bundle["text/markdown"].startswith("```{=latex}")

    _write_info(tmp_path, monkeypatch, "typst")
    typ = t.make(type="quarto")
    bundle = typ._repr_mimebundle_()
    assert bundle["text/typst"] == str(typ)
    assert bundle["text/markdown"].startswith("```{=typst}")

    _write_info(tmp_path, monkeypatch, "docx")
    assert t.make(type="quarto").tables

    _write_info(tmp_path, monkeypatch, "epub")
    assert type(t.make(type="quarto")).__name__ == "GT"


def test_make_without_type_displays_bundle():
    t = _table()
    with mock.patch("maketables.mtable.display") as disp:
        assert t.make() is None
    (bundle,) = disp.call_args.args
    assert set(bundle) == {"text/html", "text/markdown", "text/latex", "text/typst"}
    assert disp.call_args.kwargs == {"raw": True}


def test_ipython_display_hook():
    t = _table()
    with mock.patch("maketables.mtable.display") as disp:
        t._ipython_display_()
    assert disp.call_args.kwargs == {"raw": True}


def test_display_styles_flow_into_bundle():
    t = _table(typst_style={"notes_font_size": "7pt"}, tex_style={"texlocation": "b"})
    bundle = t._repr_mimebundle_()
    assert "7pt" in bundle["text/typst"]
    assert "[b]" in bundle["text/latex"]


def test_public_export():
    assert mt.MTable is MTable
